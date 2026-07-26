"""Readers for the two Excel document types.

Funding reports are tabular (header row + data rows); invoices store their
figures at fixed cell positions. Both are read with openpyxl and returned as
Polars DataFrames.
"""

from __future__ import annotations

import logging
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl
import polars as pl

log = logging.getLogger(__name__)

FUNDING_PATTERN = re.compile(r"Funding Report (\d{8})-(\d+)\.xlsx$", re.IGNORECASE)
INVOICE_PATTERN = re.compile(r"^BS .*Bill\.xlsx$", re.IGNORECASE)

# Weekly folders older than this are treated as bad data, not history.
EARLIEST_WEEK = date(2015, 1, 1)

CLAIM_SCHEMA = {
    "batch_id": pl.Utf8,
    "claim_number": pl.Utf8,
    "service_date": pl.Date,
    "gross_billable_charge": pl.Float64,
    "total_payment_amount": pl.Float64,
    "plan_type": pl.Utf8,
    "provider": pl.Utf8,
    "week": pl.Utf8,
    "source_file": pl.Utf8,
}

INVOICE_SCHEMA = {
    "week": pl.Utf8,
    "source_file": pl.Utf8,
    "invoice_date": pl.Date,
    "invoice_number": pl.Utf8,
    "carrier_claims": pl.Float64,
    "tpa_claims": pl.Float64,
    "total_claims": pl.Float64,
    "carrier_fees": pl.Float64,
    "total_amount_due": pl.Float64,
    "negative_batch": pl.Boolean,
}


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def _as_float(value: object) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    return None


def _as_str(value: object) -> str | None:
    if value is None:
        return None
    return str(value).strip() or None


def read_funding_report(path: Path, week: str) -> pl.DataFrame:
    """Parse one funding report into claim-level rows.

    The batch id is only guaranteed to appear once at the top of the data
    (row 2), so it is propagated to every row of the file.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
        col = {name: i for i, name in enumerate(header)}

        def cell(row: tuple, name: str) -> object:
            i = col.get(name)
            return row[i] if i is not None and i < len(row) else None

        batch_id: str | None = None
        records: list[dict] = []
        for row in rows:
            if all(v is None for v in row):
                continue
            batch_id = _as_str(cell(row, "batch_id")) or batch_id
            records.append(
                {
                    "batch_id": batch_id,
                    "claim_number": _as_str(cell(row, "claim_number")),
                    "service_date": _as_date(cell(row, "service_date")),
                    "gross_billable_charge": _as_float(cell(row, "gross_billable_charge")),
                    "total_payment_amount": _as_float(cell(row, "total_payment_amount")),
                    "plan_type": _as_str(cell(row, "plan_type")),
                    "provider": _as_str(cell(row, "provider")),
                    "week": week,
                    "source_file": path.name,
                }
            )
    finally:
        wb.close()
    frame = pl.DataFrame(records, schema=CLAIM_SCHEMA)
    log.debug(
        "week %s: parsed funding report %s: %d claim rows, %d batches, %s total payment",
        week, path.name, frame.height,
        frame["batch_id"].n_unique() if frame.height else 0,
        f"{frame['total_payment_amount'].sum() or 0.0:,.2f}",
    )
    return frame


def read_invoice(path: Path, week: str) -> pl.DataFrame:
    """Parse one fixed-layout invoice file into a single summary row.

    When "NEGATIVE BATCH" is flagged in A16, the total amount due shifts
    down one row to G17.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        negative_batch = "NEGATIVE BATCH" in str(ws.cell(row=16, column=1).value or "").upper()
        total_due_row = 17 if negative_batch else 16
        record = {
            "week": week,
            "source_file": path.name,
            "invoice_date": _as_date(ws.cell(row=3, column=6).value),
            "invoice_number": _as_str(ws.cell(row=3, column=7).value),
            "carrier_claims": _as_float(ws.cell(row=12, column=6).value),
            "tpa_claims": _as_float(ws.cell(row=13, column=6).value),
            "total_claims": _as_float(ws.cell(row=14, column=7).value),
            "carrier_fees": _as_float(ws.cell(row=15, column=7).value),
            "total_amount_due": _as_float(ws.cell(row=total_due_row, column=7).value),
            "negative_batch": negative_batch,
        }
    finally:
        wb.close()
    log.debug(
        "week %s: parsed invoice %s: number=%s, total claims=%s, total due=%s%s",
        week, path.name, record["invoice_number"],
        f"{record['total_claims']:,.2f}" if record["total_claims"] is not None else "—",
        f"{record['total_amount_due']:,.2f}" if record["total_amount_due"] is not None else "—",
        " (NEGATIVE BATCH)" if negative_batch else "",
    )
    return pl.DataFrame([record], schema=INVOICE_SCHEMA)


def _is_valid_week_dir(name: str) -> bool:
    """A weekly folder must be a real YYYYMMDD date in a plausible range.

    All-digit names that fail validation are logged, so a typo'd folder
    (e.g. 20241332 or 2024010) is noticed instead of silently skipped.
    """
    if len(name) != 8:
        log.warning("Skipping folder %s: expected YYYYMMDD (8 digits)", name)
        return False
    try:
        d = datetime.strptime(name, "%Y%m%d").date()
    except ValueError:
        log.warning("Skipping folder %s: not a valid calendar date", name)
        return False
    if d > date.today():
        log.warning("Skipping folder %s: week is in the future", name)
        return False
    if d < EARLIEST_WEEK:
        log.warning("Skipping folder %s: week is before %s", name, EARLIEST_WEEK)
        return False
    return True


def load_directory(input_path: Path) -> tuple[pl.DataFrame, pl.DataFrame]:
    """Walk the weekly folders and return (claims, invoices) DataFrames."""
    claim_frames: list[pl.DataFrame] = []
    invoice_frames: list[pl.DataFrame] = []

    week_dirs = sorted(
        d for d in input_path.iterdir()
        if d.is_dir() and d.name.isdigit() and _is_valid_week_dir(d.name)
    )
    if not week_dirs:
        raise FileNotFoundError(f"No weekly YYYYMMDD folders found under {input_path}")

    log.info("Found %d weekly folders under %s (%s .. %s)",
             len(week_dirs), input_path, week_dirs[0].name, week_dirs[-1].name)

    for week_dir in week_dirs:
        week = week_dir.name
        week_funding: list[pl.DataFrame] = []
        week_invoices: list[pl.DataFrame] = []
        for f in sorted(week_dir.iterdir()):
            if FUNDING_PATTERN.search(f.name):
                week_funding.append(read_funding_report(f, week))
            elif INVOICE_PATTERN.match(f.name):
                week_invoices.append(read_invoice(f, week))
            elif f.suffix.lower() == ".xlsx":
                log.warning("Unrecognized xlsx file skipped: %s", f)
        log.info(
            "week %s: %d funding reports (%d claim rows), %d invoice%s",
            week, len(week_funding), sum(fr.height for fr in week_funding),
            len(week_invoices), "" if len(week_invoices) == 1 else "s",
        )
        if not week_invoices:
            log.warning("week %s: no invoice file found", week)
        if not week_funding:
            log.warning("week %s: no funding report files found", week)
        claim_frames.extend(week_funding)
        invoice_frames.extend(week_invoices)

    claims = pl.concat(claim_frames) if claim_frames else pl.DataFrame(schema=CLAIM_SCHEMA)
    invoices = pl.concat(invoice_frames) if invoice_frames else pl.DataFrame(schema=INVOICE_SCHEMA)
    log.info(
        "Parsed %d weeks: %d claim rows from %d funding reports, %d invoices",
        len(week_dirs), claims.height, len(claim_frames), invoices.height,
    )
    return claims, invoices
