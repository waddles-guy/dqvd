"""Tests for the Excel readers in dqvd/parsing.py.

Real .xlsx files are generated with openpyxl into temp dirs, so these tests
exercise the actual workbook-reading path end to end.

Run with:  python -m unittest discover tests -v
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl
import polars as pl

from dqvd.parsing import (
    CLAIM_SCHEMA,
    EARLIEST_WEEK,
    INVOICE_SCHEMA,
    _as_date,
    _as_float,
    _as_str,
    load_directory,
    read_funding_report,
    read_invoice,
)

FUNDING_HEADER = [
    "batch_id", "claim_number", "service_date", "gross_billable_charge",
    "total_payment_amount", "plan_type", "provider",
]


def write_funding(path: Path, rows: list[list], header: list[str] = FUNDING_HEADER) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for row in rows:
        ws.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_invoice(path: Path, *, invoice_date=None, invoice_number="INV-1",
                  carrier_claims=100.0, tpa_claims=50.0, total_claims=150.0,
                  carrier_fees=10.0, total_due=160.0, a16=None,
                  negative_layout=False) -> Path:
    """Fixed-layout invoice: F3 date, G3 number, F12/F13 claims, G14 total,
    G15 fees, G16 total due (G17 when the negative-batch layout is used)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=3, column=6, value=invoice_date)
    ws.cell(row=3, column=7, value=invoice_number)
    ws.cell(row=12, column=6, value=carrier_claims)
    ws.cell(row=13, column=6, value=tpa_claims)
    ws.cell(row=14, column=7, value=total_claims)
    ws.cell(row=15, column=7, value=carrier_fees)
    if a16 is not None:
        ws.cell(row=16, column=1, value=a16)
    ws.cell(row=17 if negative_layout else 16, column=7, value=total_due)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


class CoercionTests(unittest.TestCase):
    def test_as_date(self):
        self.assertEqual(_as_date(datetime(2026, 7, 13, 10, 30)), date(2026, 7, 13))
        self.assertEqual(_as_date(date(2026, 7, 13)), date(2026, 7, 13))
        self.assertEqual(_as_date("2026-07-13"), date(2026, 7, 13))
        self.assertEqual(_as_date("  2026-07-13  "), date(2026, 7, 13))
        self.assertIsNone(_as_date("not a date"))
        self.assertIsNone(_as_date("13/07/2026"))
        self.assertIsNone(_as_date(None))
        self.assertIsNone(_as_date(20260713))

    def test_as_float(self):
        self.assertEqual(_as_float(5), 5.0)
        self.assertEqual(_as_float(5.5), 5.5)
        self.assertEqual(_as_float(-0.01), -0.01)
        self.assertEqual(_as_float(0), 0.0)
        self.assertIsNone(_as_float("5.5"))  # strings are not coerced
        self.assertIsNone(_as_float(None))

    def test_as_str(self):
        self.assertEqual(_as_str("  hello  "), "hello")
        self.assertEqual(_as_str(123), "123")
        self.assertIsNone(_as_str(None))
        self.assertIsNone(_as_str(""))
        self.assertIsNone(_as_str("   "))


class FundingReportTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def read(self, rows: list[list], header: list[str] = FUNDING_HEADER) -> pl.DataFrame:
        path = write_funding(self.tmp / "Funding Report 20260713-1.xlsx", rows, header)
        return read_funding_report(path, "20260713")

    def test_basic_rows_parse_with_week_and_source_file(self):
        out = self.read([
            ["B1", "C1", datetime(2026, 7, 1), 100.0, 80.0, "PPO", "Prov A"],
            ["B1", "C2", datetime(2026, 7, 2), 200.0, 150.0, "HMO", "Prov B"],
        ])
        self.assertEqual(out.height, 2)
        self.assertEqual(dict(out.schema), CLAIM_SCHEMA)
        row = out.row(0, named=True)
        self.assertEqual(row["batch_id"], "B1")
        self.assertEqual(row["claim_number"], "C1")
        self.assertEqual(row["service_date"], date(2026, 7, 1))
        self.assertEqual(row["gross_billable_charge"], 100.0)
        self.assertEqual(row["total_payment_amount"], 80.0)
        self.assertEqual(row["plan_type"], "PPO")
        self.assertEqual(row["provider"], "Prov A")
        self.assertEqual(row["week"], "20260713")
        self.assertEqual(row["source_file"], "Funding Report 20260713-1.xlsx")

    def test_batch_id_propagates_to_rows_where_it_is_blank(self):
        out = self.read([
            ["B1", "C1", None, 100.0, 80.0, None, None],
            [None, "C2", None, 100.0, 80.0, None, None],
            ["  ", "C3", None, 100.0, 80.0, None, None],  # whitespace counts as blank
        ])
        self.assertEqual(out["batch_id"].to_list(), ["B1", "B1", "B1"])

    def test_new_batch_id_mid_file_switches_propagation(self):
        out = self.read([
            ["B1", "C1", None, 1.0, 1.0, None, None],
            [None, "C2", None, 1.0, 1.0, None, None],
            ["B2", "C3", None, 1.0, 1.0, None, None],
            [None, "C4", None, 1.0, 1.0, None, None],
        ])
        self.assertEqual(out["batch_id"].to_list(), ["B1", "B1", "B2", "B2"])

    def test_fully_blank_rows_are_skipped(self):
        out = self.read([
            ["B1", "C1", None, 1.0, 1.0, None, None],
            [None, None, None, None, None, None, None],
            [None, "C2", None, 1.0, 1.0, None, None],
        ])
        self.assertEqual(out.height, 2)
        self.assertEqual(out["claim_number"].to_list(), ["C1", "C2"])

    def test_columns_are_mapped_by_header_name_not_position(self):
        out = self.read(
            [[80.0, "C1", "B1"]],
            header=["total_payment_amount", "claim_number", "batch_id"],
        )
        row = out.row(0, named=True)
        self.assertEqual(row["batch_id"], "B1")
        self.assertEqual(row["claim_number"], "C1")
        self.assertEqual(row["total_payment_amount"], 80.0)
        self.assertIsNone(row["provider"])  # absent column -> null

    def test_header_names_are_stripped_and_extra_columns_ignored(self):
        out = self.read(
            [["B1", "C1", 80.0, "noise"]],
            header=["  batch_id ", "claim_number", "total_payment_amount", "unknown_col"],
        )
        row = out.row(0, named=True)
        self.assertEqual(row["batch_id"], "B1")
        self.assertEqual(row["total_payment_amount"], 80.0)

    def test_string_service_date_and_non_numeric_amounts(self):
        out = self.read([
            ["B1", "C1", "2026-07-01", "not-a-number", 80.0, None, None],
        ])
        row = out.row(0, named=True)
        self.assertEqual(row["service_date"], date(2026, 7, 1))
        self.assertIsNone(row["gross_billable_charge"])

    def test_numeric_claim_number_becomes_string(self):
        out = self.read([["B1", 12345678, None, 1.0, 1.0, None, None]])
        self.assertEqual(out.row(0, named=True)["claim_number"], "12345678")

    def test_header_only_file_yields_empty_frame_with_schema(self):
        out = self.read([])
        self.assertEqual(out.height, 0)
        self.assertEqual(dict(out.schema), CLAIM_SCHEMA)


class InvoiceTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def test_normal_invoice(self):
        path = write_invoice(
            self.tmp / "BS Weekly Bill.xlsx",
            invoice_date=datetime(2026, 7, 13), invoice_number="INV-42",
            carrier_claims=100.0, tpa_claims=50.0, total_claims=150.0,
            carrier_fees=10.0, total_due=160.0,
        )
        out = read_invoice(path, "20260713")
        self.assertEqual(out.height, 1)
        self.assertEqual(dict(out.schema), INVOICE_SCHEMA)
        row = out.row(0, named=True)
        self.assertEqual(row["week"], "20260713")
        self.assertEqual(row["source_file"], "BS Weekly Bill.xlsx")
        self.assertEqual(row["invoice_date"], date(2026, 7, 13))
        self.assertEqual(row["invoice_number"], "INV-42")
        self.assertEqual(row["carrier_claims"], 100.0)
        self.assertEqual(row["tpa_claims"], 50.0)
        self.assertEqual(row["total_claims"], 150.0)
        self.assertEqual(row["carrier_fees"], 10.0)
        self.assertEqual(row["total_amount_due"], 160.0)
        self.assertFalse(row["negative_batch"])

    def test_negative_batch_shifts_total_due_to_row_17(self):
        path = write_invoice(
            self.tmp / "BS Weekly Bill.xlsx",
            a16="NEGATIVE BATCH", negative_layout=True, total_due=-75.0,
        )
        row = read_invoice(path, "20260713").row(0, named=True)
        self.assertTrue(row["negative_batch"])
        self.assertEqual(row["total_amount_due"], -75.0)

    def test_negative_batch_marker_is_case_insensitive_and_substring(self):
        path = write_invoice(
            self.tmp / "BS Weekly Bill.xlsx",
            a16="note: negative batch this week", negative_layout=True, total_due=-5.0,
        )
        row = read_invoice(path, "20260713").row(0, named=True)
        self.assertTrue(row["negative_batch"])
        self.assertEqual(row["total_amount_due"], -5.0)

    def test_other_text_in_a16_is_not_a_negative_batch(self):
        path = write_invoice(self.tmp / "BS Weekly Bill.xlsx", a16="TOTAL AMOUNT DUE")
        row = read_invoice(path, "20260713").row(0, named=True)
        self.assertFalse(row["negative_batch"])
        self.assertEqual(row["total_amount_due"], 160.0)

    def test_missing_cells_become_null(self):
        wb = openpyxl.Workbook()
        path = self.tmp / "BS Empty Bill.xlsx"
        wb.save(path)
        row = read_invoice(path, "20260713").row(0, named=True)
        self.assertIsNone(row["invoice_date"])
        self.assertIsNone(row["invoice_number"])
        self.assertIsNone(row["total_claims"])
        self.assertIsNone(row["total_amount_due"])
        self.assertFalse(row["negative_batch"])


class LoadDirectoryTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def test_walks_week_folders_and_routes_files_by_pattern(self):
        write_funding(
            self.root / "20260713" / "Funding Report 20260713-1.xlsx",
            [["B1", "C1", None, 100.0, 80.0, None, None]],
        )
        write_funding(
            self.root / "20260720" / "Funding Report 20260720-1.xlsx",
            [["B2", "C2", None, 100.0, 80.0, None, None]],
        )
        write_invoice(self.root / "20260713" / "BS Weekly Bill.xlsx")

        claims, invoices = load_directory(self.root)
        self.assertEqual(claims.height, 2)
        self.assertEqual(sorted(claims["week"].to_list()), ["20260713", "20260720"])
        self.assertEqual(invoices.height, 1)
        self.assertEqual(invoices.row(0, named=True)["week"], "20260713")

    def test_week_comes_from_folder_name_not_file_name(self):
        write_funding(
            self.root / "20260720" / "Funding Report 20260713-1.xlsx",
            [["B1", "C1", None, 1.0, 1.0, None, None]],
        )
        claims, _ = load_directory(self.root)
        self.assertEqual(claims["week"].to_list(), ["20260720"])

    def test_patterns_are_case_insensitive(self):
        write_funding(
            self.root / "20260713" / "FUNDING REPORT 20260713-2.XLSX",
            [["B1", "C1", None, 1.0, 1.0, None, None]],
        )
        write_invoice(self.root / "20260713" / "bs weekly bill.xlsx")
        claims, invoices = load_directory(self.root)
        self.assertEqual(claims.height, 1)
        self.assertEqual(invoices.height, 1)

    def test_unrecognized_xlsx_is_skipped_with_warning(self):
        write_funding(
            self.root / "20260713" / "Funding Report 20260713-1.xlsx",
            [["B1", "C1", None, 1.0, 1.0, None, None]],
        )
        write_funding(self.root / "20260713" / "random notes.xlsx", [])
        with self.assertLogs("dqvd.parsing", level="WARNING") as cm:
            claims, invoices = load_directory(self.root)
        self.assertEqual(claims.height, 1)
        self.assertEqual(invoices.height, 0)
        self.assertTrue(any("random notes.xlsx" in m for m in cm.output))

    def test_non_digit_folders_are_ignored(self):
        write_funding(
            self.root / "20260713" / "Funding Report 20260713-1.xlsx",
            [["B1", "C1", None, 1.0, 1.0, None, None]],
        )
        write_funding(
            self.root / "archive" / "Funding Report 20260101-1.xlsx",
            [["OLD", "C9", None, 1.0, 1.0, None, None]],
        )
        claims, _ = load_directory(self.root)
        self.assertEqual(claims["batch_id"].to_list(), ["B1"])

    def _write_week(self, folder: str, batch: str = "B1"):
        write_funding(
            self.root / folder / f"Funding Report {folder}-1.xlsx",
            [[batch, "C1", None, 1.0, 1.0, None, None]],
        )

    def test_digit_folder_with_wrong_length_is_skipped_with_warning(self):
        self._write_week("20260713", "GOOD")
        self._write_week("2026071", "BAD")      # 7 digits
        self._write_week("202607130", "BAD2")   # 9 digits
        with self.assertLogs("dqvd.parsing", level="WARNING") as cm:
            claims, _ = load_directory(self.root)
        self.assertEqual(claims["batch_id"].to_list(), ["GOOD"])
        self.assertTrue(any("2026071" in m and "YYYYMMDD" in m for m in cm.output))

    def test_impossible_month_or_day_is_skipped_with_warning(self):
        self._write_week("20260713", "GOOD")
        self._write_week("20261301", "BADMONTH")  # month 13
        self._write_week("20260230", "BADDAY")    # Feb 30
        self._write_week("20260700", "ZERODAY")   # day 0
        with self.assertLogs("dqvd.parsing", level="WARNING") as cm:
            claims, _ = load_directory(self.root)
        self.assertEqual(claims["batch_id"].to_list(), ["GOOD"])
        self.assertEqual(sum("not a valid calendar date" in m for m in cm.output), 3)

    def test_future_week_is_skipped_with_warning(self):
        self._write_week("20260713", "GOOD")
        future = (date.today() + timedelta(days=30)).strftime("%Y%m%d")
        self._write_week(future, "FUTURE")
        with self.assertLogs("dqvd.parsing", level="WARNING") as cm:
            claims, _ = load_directory(self.root)
        self.assertEqual(claims["batch_id"].to_list(), ["GOOD"])
        self.assertTrue(any("future" in m for m in cm.output))

    def test_today_is_a_valid_week(self):
        today = date.today().strftime("%Y%m%d")
        self._write_week(today, "TODAY")
        claims, _ = load_directory(self.root)
        self.assertEqual(claims["batch_id"].to_list(), ["TODAY"])

    def test_distant_past_week_is_skipped_with_warning(self):
        self._write_week("20260713", "GOOD")
        self._write_week("19991231", "ANCIENT")
        with self.assertLogs("dqvd.parsing", level="WARNING") as cm:
            claims, _ = load_directory(self.root)
        self.assertEqual(claims["batch_id"].to_list(), ["GOOD"])
        self.assertTrue(any("before" in m for m in cm.output))

    def test_earliest_week_boundary_is_valid(self):
        self._write_week(EARLIEST_WEEK.strftime("%Y%m%d"), "OLDEST")
        claims, _ = load_directory(self.root)
        self.assertEqual(claims["batch_id"].to_list(), ["OLDEST"])

    def test_only_invalid_week_folders_raises(self):
        self._write_week("20261301", "BAD")
        with self.assertLogs("dqvd.parsing", level="WARNING"):
            with self.assertRaises(FileNotFoundError):
                load_directory(self.root)

    def test_no_week_folders_raises(self):
        (self.root / "not-a-week").mkdir()
        with self.assertRaises(FileNotFoundError):
            load_directory(self.root)

    def test_per_week_summary_is_logged_at_info(self):
        self._write_week("20260713")
        write_invoice(self.root / "20260713" / "BS Weekly Bill.xlsx")
        with self.assertLogs("dqvd.parsing", level="INFO") as cm:
            load_directory(self.root)
        self.assertTrue(any("week 20260713: 1 funding reports (1 claim rows), 1 invoice" in m
                            for m in cm.output))
        self.assertTrue(any("Found 1 weekly folders" in m for m in cm.output))

    def test_week_without_invoice_logs_warning(self):
        self._write_week("20260713")
        with self.assertLogs("dqvd.parsing", level="WARNING") as cm:
            load_directory(self.root)
        self.assertTrue(any("week 20260713: no invoice file found" in m for m in cm.output))

    def test_week_folders_with_no_files_return_empty_frames_with_schema(self):
        (self.root / "20260713").mkdir()
        claims, invoices = load_directory(self.root)
        self.assertEqual(claims.height, 0)
        self.assertEqual(invoices.height, 0)
        self.assertEqual(dict(claims.schema), CLAIM_SCHEMA)
        self.assertEqual(dict(invoices.schema), INVOICE_SCHEMA)


if __name__ == "__main__":
    unittest.main()
